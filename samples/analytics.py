"""Bounded ingestion and deterministic analytics. No model-generated code."""
import csv
import io
import math
import zipfile
import os
import numpy as np
import pandas as pd
from openpyxl import load_workbook

MAX_COLUMNS = 10_000
MAX_ROWS = 100_000
MAX_CELLS = int(os.getenv('MAX_DATA_CELLS', '1000000'))
MAX_BYTES = int(os.getenv('MAX_DATA_BYTES', str(25 * 1024 * 1024)))


def ingest(raw: bytes, suffix: str) -> pd.DataFrame:
    if len(raw) > MAX_BYTES:
        raise ValueError(f'File exceeds the {MAX_BYTES // (1024 * 1024)} MB upload limit.')
    if suffix == '.xlsx':
        with zipfile.ZipFile(io.BytesIO(raw)) as archive:
            if sum(i.file_size for i in archive.infolist()) > int(os.getenv('MAX_EXPANDED_XLSX_BYTES', str(100 * 1024 * 1024))):
                raise ValueError('Expanded workbook exceeds this server’s supported memory limit.')
        book = load_workbook(io.BytesIO(raw), read_only=True, data_only=True, keep_links=False)
        sheet = book.worksheets[0]
        # Ignore stale worksheet dimensions; enforce limits on actual streamed cells.
        sheet.reset_dimensions()
        rows = sheet.iter_rows(values_only=True)
    elif suffix == '.csv':
        rows = csv.reader(io.StringIO(raw.decode('utf-8-sig')))
        book = None
    else:
        raise ValueError('Only .csv and .xlsx files are supported.')
    try:
        header = next(rows, None)
        if not header or len(header) > MAX_COLUMNS:
            raise ValueError('A header with 1–10,000 columns is required.')
        names = [str(v).strip() if v is not None else '' for v in header]
        if any(not n for n in names) or len(set(names)) != len(names):
            raise ValueError('Column headers must be nonempty and unique.')
        values = []
        for row in rows:
            if len(row) > len(names):
                raise ValueError('A row contains more cells than the header.')
            if len(values) >= min(MAX_ROWS, MAX_CELLS // len(names)):
                raise ValueError(f'Limit exceeded: {MAX_ROWS:,} rows or {MAX_CELLS:,} cells.')
            if any(isinstance(v, str) and len(v) > 32_767 for v in row):
                raise ValueError('A cell exceeds 32,767 characters.')
            values.append(list(row) + [None] * (len(names) - len(row)))
        if not values:
            raise ValueError('No data rows found.')
        frame = pd.DataFrame(values, columns=names).replace('', None)
        for name in names:
            series = frame[name]
            # Native XLSX date cells must retain dates, not become epoch nanoseconds.
            if pd.api.types.is_datetime64_any_dtype(series):
                continue
            numeric = pd.to_numeric(series, errors='coerce')
            if series.notna().sum() and numeric.notna().sum() == series.notna().sum():
                frame[name] = numeric.replace([np.inf, -np.inf], np.nan)
        if frame.memory_usage(deep=True).sum() > int(os.getenv('MAX_DECODED_BYTES', str(200 * 1024 * 1024))):
            raise ValueError('Decoded data exceeds this server’s supported memory limit.')
        return frame
    finally:
        if book:
            book.close()


def clean(value):
    if value is None or pd.isna(value):
        return None
    if isinstance(value, (np.integer, np.floating)):
        return value.item()
    if hasattr(value, 'isoformat'):
        return value.isoformat()
    return value


def column_info(frame, names):
    result = []
    for name in names:
        s = frame[name]
        item = {'name': name, 'type': str(s.dtype), 'missing': int(s.isna().sum())}
        if pd.api.types.is_numeric_dtype(s):
            item.update({k: clean(v) for k, v in s.describe().to_dict().items()})
        result.append(item)
    return result


def calculate(frame, column, operation, timestamp=None, unit=None, max_gap_seconds=None):
    if column not in frame or not pd.api.types.is_numeric_dtype(frame[column]):
        raise ValueError('Select a numeric column.')
    s = frame[column]
    valid = s.dropna()
    if valid.empty:
        raise ValueError('The selected column has no numeric observations.')
    result = {'operation': operation, 'column': column, 'observations': len(valid), 'missing': int(s.isna().sum())}
    if operation in ('min', 'max', 'mean', 'sum'):
        result.update(value=float(getattr(valid, operation)()), unit='source units')
    elif operation == 'trend':
        first, last = float(valid.iloc[0]), float(valid.iloc[-1])
        result.update(value=last-first, unit='source units', method='Last minus first valid observation in row order; not a rate.')
    elif operation == 'energy':
        if timestamp not in frame or unit not in ('W', 'kW'):
            raise ValueError('Energy requires a timestamp column and explicit W or kW power units.')
        if pd.api.types.is_numeric_dtype(frame[timestamp]):
            raise ValueError('Numeric timestamps are unsupported. Use ISO 8601 dates or formatted Excel date cells.')
        times = pd.to_datetime(frame[timestamp], errors='coerce', utc=True)
        if times.isna().any() or not times.is_monotonic_increasing or times.duplicated().any():
            raise ValueError('Timestamps must all be valid, unique, and strictly increasing.')
        dt = times.diff().dt.total_seconds().to_numpy()[1:]
        if not len(dt):
            raise ValueError('At least two timestamped observations are required.')
        if max_gap_seconds is None or not math.isfinite(max_gap_seconds) or max_gap_seconds <= 0:
            raise ValueError('Specify the maximum accepted interval in seconds.')
        p = s.to_numpy(dtype=float)
        accepted = np.isfinite(p[:-1]) & np.isfinite(p[1:]) & (dt <= max_gap_seconds)
        if not accepted.any():
            raise ValueError('No valid adjacent intervals within the accepted gap.')
        energy = ((p[:-1][accepted] + p[1:][accepted]) / 2 * dt[accepted]).sum() / 3600
        result.update(value=float(energy / (1000 if unit == 'W' else 1)), unit='kWh',
                      method='Trapezoidal integration of adjacent valid samples; no extrapolation.',
                      covered_seconds=float(dt[accepted].sum()), excluded_intervals=int((~accepted).sum()),
                      partial=bool((~accepted).any()), power_unit=unit, max_gap_seconds=max_gap_seconds)
    else:
        raise ValueError('Unsupported operation.')
    if not math.isfinite(result['value']):
        raise ValueError('Calculation exceeds the supported numeric range.')
    return result


def synthetic():
    rng = np.random.default_rng(42)
    t = pd.date_range('2026-06-15 06:00', periods=43_201, freq='s')
    daylight = np.maximum(0, np.sin(np.linspace(0, np.pi, len(t))))
    solar = np.maximum(0, 5200 * daylight + rng.normal(0, 55, len(t)))
    return pd.DataFrame({'Timestamp': t.astype(str), 'Solar power (W)': solar.round(2),
                         'Load power (W)': (1400 + 300 * np.sin(np.linspace(0, 16, len(t))) + rng.normal(0, 65, len(t))).round(2),
                         'Voltage (V)': rng.normal(230, 1.2, len(t)).round(2),
                         'Panel temperature (°C)': (22 + daylight * 23 + rng.normal(0, .4, len(t))).round(2)})
